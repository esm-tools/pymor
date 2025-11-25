#!/usr/bin/env python3
"""
Create Excel file with CMIP7 variables pre-populated from data request
"""

import json
import pandas as pd
from collections import defaultdict
from pathlib import Path

print("="*80)
print("CREATING CMIP7 VARIABLE MAPPING EXCEL FILE")
print("="*80)

# Load the data request metadata
print("\nLoading CMIP7 Data Request JSON files...")
with open('dreq_v1.2.2.2_metadata.json', 'r') as f:
    dreq_metadata = json.load(f)

compound_names = dreq_metadata.get('Compound Name', {})
print(f"Total compound names found: {len(compound_names)}")

# Parse compound names to extract unique variables
# Format: realm.variable.method-level-grid-type.frequency.region
variables_dict = defaultdict(lambda: {
    'realms': set(),
    'frequencies': set(),
    'regions': set(),
    'compound_names': []
})

for compound_name, details in compound_names.items():
    parts = compound_name.split('.')
    if len(parts) >= 2:
        realm = parts[0]
        variable = parts[1]
        frequency = parts[3] if len(parts) > 3 else 'unknown'
        region = parts[4] if len(parts) > 4 else 'GLB'

        variables_dict[variable]['realms'].add(realm)
        variables_dict[variable]['frequencies'].add(frequency)
        variables_dict[variable]['regions'].add(region)
        variables_dict[variable]['compound_names'].append(compound_name)

        # Store details from first occurrence
        if isinstance(details, dict) and 'details' not in variables_dict[variable]:
            variables_dict[variable]['details'] = details

print(f"Unique CMIP7 variables extracted: {len(variables_dict)}")

# Create DataFrame with all variables
print("\nCreating DataFrame with pre-populated CMIP7 metadata...")
data = []

for var_id in sorted(variables_dict.keys()):
    var_info = variables_dict[var_id]
    details = var_info.get('details', {})

    row = {
        # CMIP7 metadata (pre-populated)
        'variable_id': var_id,
        'standard_name': details.get('standard_name', '') if isinstance(details, dict) else '',
        'long_name': details.get('title', '') if isinstance(details, dict) else '',
        'units': details.get('units', '') if isinstance(details, dict) else '',
        'frequency': ', '.join(sorted(var_info['frequencies'])),
        'modeling_realm': ', '.join(sorted(var_info['realms'])),

        # Model-specific mappings (empty - for user input)
        'fesom': '',
        'oifs': '',
        'recom': '',
        'lpj_guess': '',

        # Processing information (empty - for user input)
        'preprocess': '',
        'formula': '',
        'comment': '',
        'status': 'pending',
        'priority': '',
    }
    data.append(row)

df = pd.DataFrame(data)

# Create Excel file with formatting
output_file = 'cmip7_variable_mapping.xlsx'
print(f"\nWriting to {output_file}...")

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Variable Mapping', index=False)

    # Get the worksheet
    worksheet = writer.sheets['Variable Mapping']

    # Set column widths for better readability
    column_widths = {
        'A': 20,  # variable_id
        'B': 40,  # standard_name
        'C': 50,  # long_name
        'D': 15,  # units
        'E': 20,  # frequency
        'F': 20,  # modeling_realm
        'G': 20,  # fesom
        'H': 20,  # oifs
        'I': 20,  # recom
        'J': 20,  # lpj_guess
        'K': 30,  # preprocess
        'L': 40,  # formula
        'M': 50,  # comment
        'N': 15,  # status
        'O': 15,  # priority
    }

    for col, width in column_widths.items():
        worksheet.column_dimensions[col].width = width

    # Freeze the header row
    worksheet.freeze_panes = 'A2'

    # Add data validation for status column
    from openpyxl.worksheet.datavalidation import DataValidation

    status_validation = DataValidation(
        type="list",
        formula1='"pending,in_progress,completed,not_applicable"',
        allow_blank=True
    )
    status_validation.error = 'Please select from the dropdown list'
    status_validation.errorTitle = 'Invalid Status'
    worksheet.add_data_validation(status_validation)
    status_validation.add(f'N2:N{len(df)+1}')

    # Add data validation for priority column
    priority_validation = DataValidation(
        type="list",
        formula1='"high,medium,low"',
        allow_blank=True
    )
    priority_validation.error = 'Please select from the dropdown list'
    priority_validation.errorTitle = 'Invalid Priority'
    worksheet.add_data_validation(priority_validation)
    priority_validation.add(f'O2:O{len(df)+1}')

    # Style the header row
    from openpyxl.styles import Font, PatternFill, Alignment

    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Color code different column groups
    # CMIP7 metadata columns (light blue)
    cmip7_fill = PatternFill(start_color='E7F3FF', end_color='E7F3FF', fill_type='solid')
    for row in range(2, len(df) + 2):
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            worksheet[f'{col}{row}'].fill = cmip7_fill

    # Model mapping columns (light green)
    model_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    for row in range(2, len(df) + 2):
        for col in ['G', 'H', 'I', 'J']:
            worksheet[f'{col}{row}'].fill = model_fill

    # Processing columns (light yellow)
    process_fill = PatternFill(start_color='FFF9E6', end_color='FFF9E6', fill_type='solid')
    for row in range(2, len(df) + 2):
        for col in ['K', 'L', 'M', 'N', 'O']:
            worksheet[f'{col}{row}'].fill = process_fill

print(f"\n✓ Excel file created successfully: {output_file}")
print(f"  - Total variables: {len(df)}")
print(f"  - Total columns: {len(df.columns)}")
print(f"  - CMIP7 metadata columns (blue): 6")
print(f"  - Model mapping columns (green): 4")
print(f"  - Processing columns (yellow): 5")

print("\n" + "="*80)
print("USAGE INSTRUCTIONS")
print("="*80)
print("""
1. Open cmip7_variable_mapping.xlsx in Excel or LibreOffice
2. CMIP7 metadata columns (blue) are pre-populated - DO NOT EDIT
3. Fill in model-specific variable names in green columns:
   - fesom: FESOM variable name
   - oifs: OIFS variable name
   - recom: REcoM variable name
   - lpj_guess: LPJ-Guess variable name
4. Fill in processing information in yellow columns:
   - preprocess: e.g., 'direct', 'avg24h', 'surface_extraction'
   - formula: calculation formula for derived variables
   - comment: any additional notes
   - status: select from dropdown (pending/in_progress/completed/not_applicable)
   - priority: select from dropdown (high/medium/low)
5. Save the Excel file
6. Run the conversion script to generate YAML for pycmor

Example entries:
  - tas: oifs='t2m', preprocess='daily_mean', status='completed'
  - thetao: fesom='temp', preprocess='direct', status='completed'
  - sos: fesom='salt', preprocess='surface_extraction', status='in_progress'
""")

print("="*80)
print("Next step: Use excel_to_yaml.py to convert the filled Excel to YAML")
print("="*80)
